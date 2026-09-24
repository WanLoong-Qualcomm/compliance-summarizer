"""Safe, read-only opening of supported compliance workbooks."""

from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


SUPPORTED_WORKBOOK_EXTENSIONS = frozenset({".xlsx", ".xlsm"})


class WorkbookError(ValueError):
    """Actionable error raised while validating or opening a workbook."""


@dataclass(slots=True)
class OpenedWorkbook:
    """An opened workbook and its selected worksheet.

    Instances are yielded only inside :func:`open_workbook`; the context
    manager closes the underlying workbook after the caller finishes reading.
    """

    source_path: Path
    sheet_name: str
    available_sheet_names: tuple[str, ...]
    workbook: Any
    worksheet: Any


def validate_workbook_path(path: str | Path) -> Path:
    """Validate existence and supported extension without opening the file."""

    if not isinstance(path, (str, Path)):
        raise WorkbookError("Workbook path must be a string or pathlib.Path.")
    if isinstance(path, str) and not path.strip():
        raise WorkbookError("Workbook path must be nonempty.")

    source_path = Path(path)
    if not source_path.exists():
        raise WorkbookError(
            f"Workbook file '{source_path}' does not exist. "
            "Check the configured excel_file_path."
        )
    if not source_path.is_file():
        raise WorkbookError(
            f"Workbook path '{source_path}' is not a file. "
            "Set excel_file_path to an .xlsx or .xlsm file."
        )
    if source_path.suffix.lower() not in SUPPORTED_WORKBOOK_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_WORKBOOK_EXTENSIONS))
        raise WorkbookError(
            f"Unsupported workbook extension for '{source_path}'. "
            f"Use one of: {supported}."
        )
    return source_path


@contextmanager
def open_workbook(path: str | Path, sheet_name: str) -> Iterator[OpenedWorkbook]:
    """Open a workbook read-only and yield one exact named worksheet.

    Macros are never retained or executed. Formula cells are read from cached
    values where available. The workbook is closed on successful completion,
    sheet-selection failure, load failure, or an exception from the caller.
    """

    source_path = validate_workbook_path(path)
    if not isinstance(sheet_name, str) or not sheet_name.strip():
        raise WorkbookError("Compliance sheet name must be a nonempty string.")

    workbook: Any = None
    try:
        workbook = load_workbook(
            filename=source_path,
            read_only=True,
            data_only=True,
            keep_vba=False,
            keep_links=False,
        )
        available_sheet_names = tuple(workbook.sheetnames)
        if sheet_name not in available_sheet_names:
            available = ", ".join(repr(name) for name in available_sheet_names) or "none"
            raise WorkbookError(
                f"Compliance sheet '{sheet_name}' was not found in '{source_path}'. "
                f"Available sheets: {available}."
            )
        worksheet = workbook[sheet_name]
    except WorkbookError:
        if workbook is not None:
            _close_quietly(workbook)
        raise
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as error:
        if workbook is not None:
            _close_quietly(workbook)
        raise WorkbookError(
            f"Could not open workbook '{source_path}' read-only: {error}. "
            "Confirm that the file is a valid .xlsx or .xlsm workbook."
        ) from error
    except Exception as error:
        if workbook is not None:
            _close_quietly(workbook)
        raise WorkbookError(
            f"Could not open workbook '{source_path}' read-only: {error}. "
            "Confirm that the file is a valid .xlsx or .xlsm workbook."
        ) from error

    handle = OpenedWorkbook(
        source_path=source_path,
        sheet_name=sheet_name,
        available_sheet_names=available_sheet_names,
        workbook=workbook,
        worksheet=worksheet,
    )
    try:
        yield handle
    finally:
        _close_quietly(workbook)


def _close_quietly(workbook: Any) -> None:
    try:
        workbook.close()
    except Exception:
        pass
