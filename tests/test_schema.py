from pathlib import Path

import openpyxl

from compliance_summarizer.schema import discover_sheet_schema
from compliance_summarizer.workbook import open_workbook


FIXED_HEADERS = [
    "LNAMODE",
    "CAMODE",
    "STD",
    "BAND",
    "BW",
    "MEASPORT",
    "DLP",
    "DIV",
    "F0_MHZ",
    "TESTNAME",
    "GAINMODE",
    "BBPATH",
    "FREQ",
    "CHANNEL",
    "Result?",
    "LL",
    "UL",
]
PIVOT_STATS = ("MIN", "MAX", "NN_25C AVG", "wcMargin", "wcValue")


def create_schema_workbook(
    path: Path,
    *,
    merged: bool,
    missing_statistic: str | None = None,
    duplicate_statistic: str | None = None,
    duplicate_fixed_header: bool = False,
    trailing_columns: bool = False,
    data_row: bool = True,
) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Combined"
    for column, header in enumerate(FIXED_HEADERS, start=1):
        worksheet.cell(row=4, column=column, value=header)

    pivot_names = ("GF-PROTO", "GF-QMOM", "SEC-DR5")
    start_column = len(FIXED_HEADERS) + 1
    for pivot_index, pivot_name in enumerate(pivot_names):
        group_start = start_column + pivot_index * len(PIVOT_STATS)
        if merged:
            worksheet.cell(row=2, column=group_start, value=pivot_name)
            worksheet.merge_cells(
                start_row=2,
                start_column=group_start,
                end_row=2,
                end_column=group_start + len(PIVOT_STATS) - 1,
            )
        else:
            worksheet.cell(row=2, column=group_start, value=pivot_name)
        for statistic_index, statistic in enumerate(PIVOT_STATS):
            statistic_column = group_start + statistic_index
            value = statistic
            if statistic == missing_statistic:
                value = None
            if duplicate_statistic is not None and statistic_index == 1:
                value = duplicate_statistic
            worksheet.cell(row=3, column=statistic_column, value=value)
            worksheet.cell(row=4, column=statistic_column, value=f"Column{statistic_index + 1}")

    trailing_start = start_column + len(pivot_names) * len(PIVOT_STATS)
    if trailing_columns:
        worksheet.cell(row=3, column=trailing_start, value="DELTA")
        worksheet.cell(row=4, column=trailing_start, value="Delta (reference only)")
        worksheet.cell(row=4, column=trailing_start + 1, value="FAIL TYPE")

    if duplicate_fixed_header:
        worksheet.cell(row=4, column=trailing_start + 2, value="TESTNAME")
    if data_row:
        worksheet.cell(row=5, column=1, value="LNA")

    workbook.save(path)


def discover(path: Path, selected_pivot: str | None = "GF-PROTO"):
    with open_workbook(path, "Combined") as opened:
        return discover_sheet_schema(opened.worksheet, selected_pivot)


def test_merged_headers_discover_ordered_pivots_and_statistics(tmp_path):
    workbook_path = tmp_path / "merged.xlsx"
    create_schema_workbook(workbook_path, merged=True, trailing_columns=True)

    result = discover(workbook_path)

    assert result.is_valid
    assert result.schema is not None
    assert result.schema.pivot_names == ("GF-PROTO", "GF-QMOM", "SEC-DR5")
    assert result.schema.pivots[0].statistics == {
        "MIN": 18,
        "MAX": 19,
        "NN_25C AVG": 20,
        "wcMargin": 21,
        "wcValue": 22,
    }
    assert "Column1" not in result.schema.pivots[0].statistics
    assert "Delta (reference only)" in result.schema.unknown_columns
    assert any(finding.code == "unknown_trailing_columns" for finding in result.findings)


def test_unmerged_equivalent_headers_are_discovered(tmp_path):
    workbook_path = tmp_path / "unmerged.xlsx"
    create_schema_workbook(workbook_path, merged=False)

    result = discover(workbook_path, selected_pivot="GF-QMOM")

    assert result.is_valid
    assert result.schema is not None
    assert result.schema.pivot_names == ("GF-PROTO", "GF-QMOM", "SEC-DR5")


def test_missing_statistic_blocks_schema(tmp_path):
    workbook_path = tmp_path / "missing.xlsx"
    create_schema_workbook(workbook_path, merged=False, missing_statistic="wcMargin")

    result = discover(workbook_path)

    assert result.schema is None
    assert any(finding.code == "missing_pivot_statistics" for finding in result.blocking_findings)


def test_duplicate_statistic_and_fixed_header_block_schema(tmp_path):
    duplicate_stat_path = tmp_path / "duplicate_stat.xlsx"
    create_schema_workbook(
        duplicate_stat_path,
        merged=False,
        duplicate_statistic="MIN",
    )
    result = discover(duplicate_stat_path)
    assert result.schema is None
    assert any(finding.code == "duplicate_pivot_statistic" for finding in result.blocking_findings)

    duplicate_header_path = tmp_path / "duplicate_header.xlsx"
    create_schema_workbook(duplicate_header_path, merged=False, duplicate_fixed_header=True)
    result = discover(duplicate_header_path)
    assert result.schema is None
    assert any(finding.code == "duplicate_required_header" for finding in result.blocking_findings)


def test_selected_pivot_and_data_region_are_required(tmp_path):
    no_data_path = tmp_path / "no_data.xlsx"
    create_schema_workbook(no_data_path, merged=False, data_row=False)
    result = discover(no_data_path, selected_pivot="NOT_PRESENT")

    assert result.schema is None
    codes = {finding.code for finding in result.blocking_findings}
    assert "selected_pivot_absent" in codes
    assert "no_data_rows" in codes
