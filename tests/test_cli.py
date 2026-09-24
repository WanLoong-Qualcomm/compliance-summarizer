import json
from pathlib import Path

import openpyxl

from compliance_summarizer.cli import main


FIXED_HEADERS = (
    "LNAMODE", "CAMODE", "STD", "BAND", "BW", "MEASPORT", "DLP", "DIV",
    "F0_MHZ", "TESTNAME", "GAINMODE", "BBPATH", "FREQ", "CHANNEL", "Result?",
    "LL", "UL",
)


def create_cli_workbook(path: Path, *, negative_min: bool = False) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Combined"
    for column, header in enumerate(FIXED_HEADERS, start=1):
        worksheet.cell(row=4, column=column, value=header)
    for pivot_index, pivot_name in enumerate(("FIRST", "SECOND")):
        start = 18 + pivot_index * 5
        worksheet.cell(row=2, column=start, value=pivot_name)
        for offset, statistic in enumerate(("MIN", "MAX", "NN_25C AVG", "wcMargin", "wcValue")):
            worksheet.cell(row=3, column=start + offset, value=statistic)
            worksheet.cell(row=4, column=start + offset, value=f"Column{offset + 1}")
    values = {
        "LNAMODE": "LNA0", "CAMODE": "CA0", "STD": "5G", "BAND": "n78",
        "BW": 100, "MEASPORT": "PORT0", "DLP": "DLP0", "DIV": 2,
        "F0_MHZ": 3600, "TESTNAME": "GAIN", "GAINMODE": 1, "BBPATH": "BB0",
        "FREQ": 3600, "CHANNEL": 1, "Result?": "PASS", "LL": -1, "UL": 1,
    }
    for header, value in values.items():
        worksheet.cell(row=5, column=FIXED_HEADERS.index(header) + 1, value=value)
    for pivot_index in range(2):
        start = 18 + pivot_index * 5
        minimum = -0.5 if negative_min and pivot_index == 0 else 0.5
        for offset, value in enumerate((minimum, 1.0, 10.0 + pivot_index, 0.5, 10.0)):
            worksheet.cell(row=5, column=start + offset, value=value)
    workbook.save(path)


def test_cli_rejects_uncompleted_new_settings(capsys, monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr("builtins.input", lambda _prompt: "")

    assert main([]) == 2

    captured = capsys.readouterr()
    assert "Settings error: Invalid settings field 'excel_file_path'" in captured.err
    assert (tmp_path / "settings.json").is_file()


def test_cli_loads_existing_valid_settings(capsys, monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr("builtins.input", lambda _prompt: "")
    workbook_path = tmp_path / "input.xlsx"
    create_cli_workbook(workbook_path)
    (tmp_path / "settings.json").write_text(
        json.dumps(
            {
                "excel_file_path": str(workbook_path),
                "compliance_sheet_name": "Combined",
                "tests_to_report": ["GAIN"],
                "acceptable_variation": {"GAIN": 0.2},
                "background_information": "",
                "pivot_field_of_interest": "FIRST",
                "bypass_model": True,
            }
        ),
        encoding="utf-8",
    )

    responses = iter(("", "", ""))
    monkeypatch.setattr("builtins.input", lambda _prompt: next(responses))

    assert main([]) == 0
    captured = capsys.readouterr()
    assert "Settings loaded successfully." in captured.out
    assert "Model stage: AI generation is bypassed in v0.1." in captured.out
    assert "GAIN rows: 1" in captured.out
    assert (tmp_path / "compliance-summary.html").is_file()


def test_cli_reports_invalid_existing_settings_without_stack_trace(
    capsys, monkeypatch, tmp_path
):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr("builtins.input", lambda _prompt: "")
    (tmp_path / "settings.json").write_text("{}", encoding="utf-8")

    assert main([]) == 2
    captured = capsys.readouterr()
    assert "Settings error:" in captured.err
    assert "Traceback" not in captured.err


def test_cli_rejects_negative_gain_before_writing_report(
    capsys, monkeypatch, tmp_path
):
    monkeypatch.chdir(tmp_path)
    workbook_path = tmp_path / "input.xlsx"
    create_cli_workbook(workbook_path, negative_min=True)
    (tmp_path / "settings.json").write_text(
        json.dumps(
            {
                "excel_file_path": str(workbook_path),
                "compliance_sheet_name": "Combined",
                "tests_to_report": ["GAIN"],
                "acceptable_variation": {"GAIN": 0.2},
                "background_information": "",
                "pivot_field_of_interest": "FIRST",
                "bypass_model": True,
            }
        ),
        encoding="utf-8",
    )
    responses = iter(("", "", "", "no"))
    monkeypatch.setattr("builtins.input", lambda _prompt: next(responses))

    assert main([]) == 2
    captured = capsys.readouterr()
    assert "Negative-GAIN check rejected" in captured.out
    assert "Report written to:" not in captured.out
    assert not (tmp_path / "compliance-summary.html").exists()


def test_cli_aborts_before_completion_when_settings_pause_reaches_eof(
    capsys, monkeypatch, tmp_path
):
    monkeypatch.chdir(tmp_path)

    def raise_eof(_prompt):
        raise EOFError

    monkeypatch.setattr("builtins.input", raise_eof)

    assert main([]) == 2
    captured = capsys.readouterr()
    assert "Input aborted:" in captured.err
    assert "Model stage:" not in captured.out
