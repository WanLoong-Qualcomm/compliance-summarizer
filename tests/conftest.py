from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook


FIXED_HEADERS = (
    "LNAMODE",
    "CAMODE",
    "STD",
    "BAND",
    "BW",
    "MEASPORT",
    "DLP",
    "DIV",
    "F0_MHZ",
    "TESTNAME",
    "GAINMODE",
    "BBPATH",
    "FREQ",
    "CHANNEL",
    "Result?",
    "LL",
    "UL",
)
PIVOT_STATS = ("MIN", "MAX", "NN_25c AVG", "wcMargin", "wcValue")


@pytest.fixture
def workbook_factory(tmp_path: Path):
    def create(
        rows: list[dict[str, object]],
        *,
        pivots: tuple[str, ...] = ("DUT-1_VAR1", "DUT-2_VAR1"),
        filename: str = "input.xlsx",
        omit_header: str | None = None,
        extra_headers: tuple[str, ...] = (),
        pivot_stats: tuple[str, ...] = PIVOT_STATS,
    ) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Combined"
        sheet.cell(1, 1, "COMPLIANCE_DASHBOARD")
        headers = [
            header
            for header in (*FIXED_HEADERS, *extra_headers)
            if header != omit_header
        ]
        for column, header in enumerate(headers, start=1):
            sheet.cell(4, column, header)
        start = len(headers) + 1
        for pivot_index, pivot in enumerate(pivots):
            first = start + pivot_index * len(pivot_stats)
            last = first + len(pivot_stats) - 1
            sheet.merge_cells(start_row=2, start_column=first, end_row=2, end_column=last)
            sheet.cell(2, first, pivot)
            for offset, statistic in enumerate(pivot_stats):
                sheet.cell(3, first + offset, statistic)
                sheet.cell(4, first + offset, f"Column{first + offset}")

        for row_number, row_data in enumerate(rows, start=5):
            for column, header in enumerate(headers, start=1):
                sheet.cell(row_number, column, row_data.get(header))
            for pivot_index, pivot in enumerate(pivots):
                first = start + pivot_index * len(pivot_stats)
                values = row_data.get(pivot, {})
                for offset, statistic in enumerate(pivot_stats):
                    sheet.cell(row_number, first + offset, values.get(statistic))
        path = tmp_path / filename
        workbook.save(path)
        workbook.close()
        return path

    return create


@pytest.fixture
def sample_rows():
    base = {
        "LNAMODE": "LNA0",
        "CAMODE": "CA0",
        "STD": "NR",
        "BAND": "n1",
        "BW": 20,
        "MEASPORT": "P1",
        "DLP": "D0",
        "DIV": 1,
        "F0_MHZ": 2100,
        "TESTNAME": "GAIN",
        "GAINMODE": 0,
        "BBPATH": "B0",
        "FREQ": 2110,
        "CHANNEL": 1,
        "Result?": "PASS",
        "LL": 8,
        "UL": 12,
    }
    return [
        {
            **base,
            "CHANNEL": 1,
            "DUT-1_VAR1": {
                "MIN": 9.2,
                "MAX": 10.4,
                "NN_25c AVG": 10.0,
                "wcMargin": -1.0,
                "wcValue": 9.2,
            },
            "DUT-2_VAR1": {
                "MIN": 8.8,
                "MAX": 10.0,
                "NN_25c AVG": 9.5,
                "wcMargin": 1.0,
                "wcValue": 9.0,
            },
        },
        {
            **base,
            "CHANNEL": 2,
            "Result?": "FAIL",
            "DUT-1_VAR1": {
                "MIN": 8.5,
                "MAX": 9.4,
                "NN_25c AVG": 9.0,
                "wcMargin": 0.0,
                "wcValue": 9.0,
            },
            "DUT-2_VAR1": {
                "MIN": 7.5,
                "MAX": 8.5,
                "NN_25c AVG": None,
                "wcMargin": -2.0,
                "wcValue": 8.0,
            },
        },
        {
            **base,
            "CHANNEL": 3,
            "DUT-1_VAR1": {
                "MIN": 8.7,
                "MAX": 9.1,
                "NN_25c AVG": "bad",
                "wcMargin": -0.2,
                "wcValue": 8.7,
            },
            "DUT-2_VAR1": {
                "MIN": 8.4,
                "MAX": 9.0,
                "NN_25c AVG": 8.8,
                "wcMargin": None,
                "wcValue": 8.4,
            },
        },
        {
            **base,
            "CHANNEL": 4,
            "TESTNAME": "SSNF",
            "DUT-1_VAR1": {"NN_25c AVG": 1.0, "wcMargin": -99},
            "DUT-2_VAR1": {"NN_25c AVG": 1.0, "wcMargin": -99},
        },
    ]


def write_settings(
    path: Path,
    workbook: Path,
    *,
    background: str = "",
    main_pivot: str = "DUT-1_VAR1",
    **overrides: object,
) -> Path:
    payload: dict[str, object] = {
        "excel_file_path": workbook.name,
        "compliance_sheet_name": "Combined",
        "block": "SIGPATH",
        "testnames": ["GAIN"],
        "background_information": background,
        "main_pivot": main_pivot,
        "group_by": [],
        "bypass_model": True,
    }
    payload.update(overrides)
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path
