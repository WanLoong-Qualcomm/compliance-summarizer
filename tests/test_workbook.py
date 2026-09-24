from pathlib import Path
from shutil import copyfile

import openpyxl
import pytest

import compliance_summarizer.workbook as workbook_module
from compliance_summarizer.workbook import WorkbookError, open_workbook


def create_workbook(path: Path, *sheet_names: str) -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = sheet_names[0]
    for sheet_name in sheet_names[1:]:
        workbook.create_sheet(sheet_name)
    workbook.save(path)


def test_supported_xlsx_and_xlsm_files_select_exact_sheet(tmp_path):
    xlsx_path = tmp_path / "compliance.xlsx"
    create_workbook(xlsx_path, "Combined", "Other")
    xlsm_path = tmp_path / "compliance.xlsm"
    copyfile(xlsx_path, xlsm_path)

    with open_workbook(xlsx_path, "Combined") as opened_xlsx:
        assert opened_xlsx.sheet_name == "Combined"
        assert opened_xlsx.worksheet.title == "Combined"
        assert opened_xlsx.available_sheet_names == ("Combined", "Other")

    with open_workbook(xlsm_path, "Combined") as opened_xlsm:
        assert opened_xlsm.worksheet.title == "Combined"


def test_missing_unsupported_and_directory_paths_fail_clearly(tmp_path):
    with pytest.raises(WorkbookError, match="does not exist"):
        open_workbook(tmp_path / "missing.xlsx", "Combined").__enter__()

    unsupported = tmp_path / "compliance.csv"
    unsupported.write_text("not a workbook", encoding="utf-8")
    with pytest.raises(WorkbookError, match="Unsupported workbook extension"):
        open_workbook(unsupported, "Combined").__enter__()

    with pytest.raises(WorkbookError, match="not a file"):
        open_workbook(tmp_path, "Combined").__enter__()


def test_corrupt_workbook_and_missing_sheet_report_actionable_errors(tmp_path):
    corrupt_path = tmp_path / "corrupt.xlsx"
    corrupt_path.write_bytes(b"not an xlsx file")
    with pytest.raises(WorkbookError, match="Could not open workbook"):
        open_workbook(corrupt_path, "Combined").__enter__()

    workbook_path = tmp_path / "compliance.xlsx"
    create_workbook(workbook_path, "Available")
    with pytest.raises(WorkbookError, match="Available sheets: 'Available'"):
        open_workbook(workbook_path, "Combined").__enter__()


def test_source_timestamp_and_contents_are_unchanged(tmp_path):
    workbook_path = tmp_path / "compliance.xlsx"
    create_workbook(workbook_path, "Combined")
    before_bytes = workbook_path.read_bytes()
    before_timestamp = workbook_path.stat().st_mtime_ns

    with open_workbook(workbook_path, "Combined") as opened:
        assert opened.worksheet.max_row == 1

    assert workbook_path.read_bytes() == before_bytes
    assert workbook_path.stat().st_mtime_ns == before_timestamp


def test_loader_uses_read_only_cached_values_and_disables_vba(monkeypatch, tmp_path):
    workbook_path = tmp_path / "compliance.xlsm"
    workbook_path.write_bytes(b"placeholder")
    calls = {}

    class FakeWorksheet:
        title = "Combined"

    class FakeWorkbook:
        sheetnames = ["Combined"]

        def __getitem__(self, name):
            assert name == "Combined"
            return FakeWorksheet()

        def close(self):
            calls["closed"] = True

    def fake_load_workbook(**kwargs):
        calls["kwargs"] = kwargs
        return FakeWorkbook()

    monkeypatch.setattr(workbook_module, "load_workbook", fake_load_workbook)
    with open_workbook(workbook_path, "Combined"):
        pass

    assert calls["kwargs"] == {
        "filename": workbook_path,
        "read_only": True,
        "data_only": True,
        "keep_vba": False,
        "keep_links": False,
    }
    assert calls["closed"] is True


def test_workbook_closes_when_sheet_selection_or_consumer_fails(monkeypatch, tmp_path):
    workbook_path = tmp_path / "compliance.xlsx"
    workbook_path.write_bytes(b"placeholder")
    closed = []

    class FakeWorkbook:
        sheetnames = ["Available"]

        def __getitem__(self, name):
            raise AssertionError("the unavailable sheet must be rejected first")

        def close(self):
            closed.append(True)

    monkeypatch.setattr(workbook_module, "load_workbook", lambda **kwargs: FakeWorkbook())
    with pytest.raises(WorkbookError):
        with open_workbook(workbook_path, "Missing"):
            pass
    assert len(closed) == 1

    class GoodFakeWorkbook(FakeWorkbook):
        sheetnames = ["Combined"]

        def __getitem__(self, name):
            return object()

    monkeypatch.setattr(workbook_module, "load_workbook", lambda **kwargs: GoodFakeWorkbook())
    with pytest.raises(RuntimeError):
        with open_workbook(workbook_path, "Combined"):
            raise RuntimeError("consumer failure")
    assert len(closed) == 2
