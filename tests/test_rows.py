from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl

from compliance_summarizer.contracts import SourceResult
from compliance_summarizer.rows import parse_data_rows
from compliance_summarizer.schema import discover_sheet_schema
from compliance_summarizer.workbook import open_workbook


FIXED_HEADERS = [
    "LNAMODE",
    "CAMODE",
    "STD",
    "BAND",
    "BW",
    "MEASPORT",
    "DLP",
    "DIV",
    "F0_MHZ",
    "TESTNAME",
    "GAINMODE",
    "BBPATH",
    "FREQ",
    "CHANNEL",
    "Result?",
    "LL",
    "UL",
]
PIVOT_STATS = ("MIN", "MAX", "NN_25C AVG", "wcMargin", "wcValue")


def create_row_workbook(path: Path, rows: list[dict[str, object]]) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Combined"
    for column, header in enumerate(FIXED_HEADERS, start=1):
        worksheet.cell(row=4, column=column, value=header)
    worksheet.cell(row=2, column=18, value="GF-PROTO")
    worksheet.cell(row=3, column=18, value="MIN")
    worksheet.cell(row=3, column=19, value="MAX")
    worksheet.cell(row=3, column=20, value="NN_25C AVG")
    worksheet.cell(row=3, column=21, value="wcMargin")
    worksheet.cell(row=3, column=22, value="wcValue")
    for column in range(18, 23):
        worksheet.cell(row=4, column=column, value=f"Column{column - 17}")

    for row_number, values in enumerate(rows, start=5):
        for header, value in values.items():
            if header in FIXED_HEADERS:
                column = FIXED_HEADERS.index(header) + 1
            else:
                column = 18 + PIVOT_STATS.index(header)
            worksheet.cell(row=row_number, column=column, value=value)
    workbook.save(path)


def parsed_rows(path: Path):
    with open_workbook(path, "Combined") as opened:
        discovery = discover_sheet_schema(opened.worksheet, selected_pivot="GF-PROTO")
        assert discovery.schema is not None
        return parse_data_rows(opened.worksheet, discovery.schema)


def inject_cached_formula(path: Path, cell_reference: str, formula: str, cached_value: str) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(path, "r") as source, NamedTemporaryFile(
        dir=path.parent,
        suffix=".xlsx",
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
        with ZipFile(temporary, "w", compression=ZIP_DEFLATED) as destination:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ElementTree.fromstring(data)
                    for cell in root.iter(f"{{{namespace}}}c"):
                        if cell.attrib.get("r") != cell_reference:
                            continue
                        cell.attrib.pop("t", None)
                        for child in list(cell):
                            cell.remove(child)
                        formula_element = ElementTree.SubElement(
                            cell, f"{{{namespace}}}f"
                        )
                        formula_element.text = formula
                        value_element = ElementTree.SubElement(cell, f"{{{namespace}}}v")
                        value_element.text = cached_value
                    data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
                destination.writestr(item, data)
    temporary_path.replace(path)


def test_complete_and_partial_rows_preserve_context_missing_values_and_zero(tmp_path):
    workbook_path = tmp_path / "rows.xlsx"
    create_row_workbook(
        workbook_path,
        [
            {
                "LNAMODE": " LNA1 ",
                "TESTNAME": " GAIN ",
                "Result?": " pass ",
                "LL": -2.0,
                "UL": 2.0,
                "MIN": -1.0,
                "MAX": 3.0,
                "NN_25C AVG": 10.0,
                "wcMargin": 0.0,
                "wcValue": 10.0,
            },
            {
                "TESTNAME": "GAIN",
                "Result?": "FAIL",
                "LL": "not numeric",
                "UL": 2.0,
                "MIN": None,
                "MAX": None,
                "NN_25C AVG": None,
                "wcMargin": "missing text",
            },
        ],
    )

    result = parsed_rows(workbook_path)

    assert [case.worksheet_row_number for case in result.cases] == [5, 6]
    first, second = result.cases
    assert first.fixed_values["LNAMODE"] == "LNA1"
    assert first.fixed_values["TESTNAME"] == "GAIN"
    assert first.source_result is SourceResult.PASS
    assert first.pivot_values["GF-PROTO"]["wcMargin"] == 0.0
    assert second.pivot_values["GF-PROTO"]["NN_25C AVG"] is None
    assert second.pivot_values["GF-PROTO"]["wcMargin"] == "missing text"
    assert second.source_values["GF-PROTO.wcMargin"] == "missing text"
    assert {finding.code for finding in result.findings} == {
        "missing_fixed_value",
        "missing_pivot_value",
        "invalid_numeric_value",
    }
    assert any(
        finding.field_name == "LL" and finding.code == "invalid_numeric_value"
        for finding in result.findings
    )


def test_formula_cached_values_are_read_as_values(tmp_path):
    workbook_path = tmp_path / "cached.xlsx"
    create_row_workbook(
        workbook_path,
        [
            {
                "TESTNAME": "GAIN",
                "Result?": "PASS",
                "LL": -1.0,
                "UL": 1.0,
                "MIN": 1.0,
                "MAX": 2.0,
                "NN_25C AVG": 10.0,
                "wcMargin": 1.0,
            }
        ],
    )
    inject_cached_formula(workbook_path, "U5", "1+1", "7")

    with open_workbook(workbook_path, "Combined") as opened:
        discovery = discover_sheet_schema(opened.worksheet, selected_pivot="GF-PROTO")
        assert discovery.schema is not None
        before = parse_data_rows(opened.worksheet, discovery.schema)

    assert before.cases[0].pivot_values["GF-PROTO"]["wcMargin"] == 7


def test_empty_data_region_returns_empty_result(tmp_path):
    workbook_path = tmp_path / "empty.xlsx"
    create_row_workbook(workbook_path, [])

    with open_workbook(workbook_path, "Combined") as opened:
        discovery = discover_sheet_schema(opened.worksheet, selected_pivot="GF-PROTO")
        assert discovery.schema is None
